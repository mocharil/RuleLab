from io import BytesIO
import asyncio
from docx import Document
import pytest
from app.services import drafting,social,ai
from app.models import SocialOptions
from app.config import settings

@pytest.mark.parametrize('kind',list(drafting.TEMPLATES))
@pytest.mark.parametrize('form',['new','amendment'])
def test_all_source_linked_templates_round_trip(client,kind,form):
    request={'title':'Pelayanan yang Inklusif','regulation_type':kind,'form':form,'region':'Bandung','institution':'Dalam Negeri','local_level':'kota','include_explanation':True,'include_attachment':True}
    preview=client.post('/api/drafting/preview',json=request)
    assert preview.status_code==200
    result=preview.json();draft=result['draft']
    assert 'Pelayanan yang Inklusif'.upper() in draft['document_heading']
    assert 'Menimbang:' in draft['preamble'] and 'Mengingat:' in draft['preamble'] and 'MEMUTUSKAN:' in draft['preamble']
    assert 'Diundangkan di' in draft['closing']
    assert '[NOMOR PENGUNDANGAN]' in draft['closing']
    assert result['sources'] and all(s['url'].startswith('https://') for s in result['sources'])
    numbers=[a['number'] for a in draft['articles']]
    assert numbers==['I','II'] if form=='amendment' else numbers==['1','2','3']
    created=client.post('/api/drafting/create',json=request)
    assert created.status_code==201
    saved=created.json()['draft']
    exported=client.get('/api/drafts/'+saved['id']+'/export')
    text='\n'.join(p.text for p in Document(BytesIO(exported.content)).paragraphs)
    assert 'Pasal '+numbers[0] in text
    assert draft['document_heading'].splitlines()[1] in text
    assert 'PENJELASAN' in text and 'LAMPIRAN' in text and '[NAMA PEJABAT PENETAP]' in text
    saved['title']='Judul diperbaiki';saved['expected_revision']=saved['revision']
    updated=client.put('/api/drafts/'+saved['id'],json=saved).json()
    assert updated['closing']==draft['closing'] and updated['drafting']['form']==form


def test_agreement_blocks_and_promulgation_media_are_distinct():
    def make(kind,**kwargs):return drafting.build(drafting.TemplateRequest(title='Layanan warga',regulation_type=kind,region='Bandung',**kwargs))['draft']
    assert 'DEWAN PERWAKILAN RAKYAT REPUBLIK' in make('UU')['preamble']
    assert 'Persetujuan Bersama' not in make('PERPPU')['preamble']
    assert 'Pasal 22 ayat (1)' in make('PERPPU')['preamble']
    city=make('PERDA_KAB',local_level='kota')
    assert 'WALI KOTA BANDUNG' in city['preamble'] and 'DPRD' not in city['preamble'] # formal name is spelled out
    assert 'LEMBARAN DAERAH KOTA BANDUNG' in city['closing']
    assert 'Persetujuan Bersama' not in make('PERWALI')['preamble']
    assert 'BERITA DAERAH KOTA BANDUNG' in make('PERWALI')['closing']
    assert 'Kesepakatan Bersama' in make('PERDES')['preamble']
    assert 'LEMBARAN DESA BANDUNG' in make('PERDES')['closing']
    assert 'BERITA DESA BANDUNG' in make('PERKADES')['closing']
    assert 'Pasal 5 ayat (1)' in make('UU',legislative_origin='PRESIDEN')['preamble']


def test_builder_never_invents_sector_specific_legal_bases_or_official_numbers():
    result=drafting.build(drafting.TemplateRequest(title='Kebijakan kesehatan',regulation_type='PERMEN'))['draft']
    assert '[BELUM DITETAPKAN]' in result['document_heading']
    assert '[Cantumkan dasar kewenangan' in result['preamble']
    assert 'Nomor 25 Tahun 2009' not in result['preamble']

def test_export_accepts_user_added_blank_lines_in_heading_and_chapters(client):
    draft=client.post('/api/drafting/create',json={'title':'Pelayanan publik'}).json()['draft']
    draft['document_heading']='\nRANCANGAN\n\nPERATURAN PEMERINTAH\n'
    draft['articles'][0]['chapter']='\nBAB I\n\nKETENTUAN UMUM\n'
    draft['expected_revision']=draft['revision']
    assert client.put('/api/drafts/'+draft['id'],json=draft).status_code==200
    result=client.get('/api/drafts/'+draft['id']+'/export')
    assert result.status_code==200
    text='\n'.join(p.text for p in Document(BytesIO(result.content)).paragraphs)
    assert 'BAB I' in text and 'KETENTUAN UMUM' in text

@pytest.mark.parametrize('count',[8,16,24,32])
def test_larger_persona_pools_have_distinct_ids_and_relevant_roles(count):
    for domain_title in ['Upah buruh','Ikan nelayan','Lingkungan sampah','Pelayanan kesehatan','Layanan digital','Tata kelola daerah']:
        people=social.get_personas_for_draft({'title':domain_title},count)
        assert len(people)==count
        assert len({p['id'] for p in people})==count
        assert all(p['synthetic'] and p['goal'] and p['cluster'] for p in people)

@pytest.mark.parametrize('count,rounds',[(8,2),(24,4),(32,5)])
def test_configured_large_simulation_contains_real_reactions_and_valid_reply_edges(client,count,rounds):
    draft=client.get('/api/drafts').json()[0]
    result=asyncio.run(social.simulate(draft,'Pasal 2\nLayanan mempunyai kanal pendampingan.',{},[],lambda *args:None,{'persona_count':count,'rounds':rounds}))
    assert result['calls_planned']==count*rounds*2
    assert result['rounds']==rounds
    for scenario in ['A','B']:
        comments=result['scenarios'][scenario]['comments']
        assert len(comments)==count*rounds
        by_id={c['id']:c for c in comments}
        for c in comments:
            if c['round']==1:assert c['reply_to'] is None
            else:
                target=by_id[c['reply_to']]
                assert target['round']<c['round'] and target['persona_id']!=c['persona_id'] and target['scenario']==scenario


def test_live_network_caps_concurrency_and_context(client,monkeypatch):
    monkeypatch.setattr(settings,'provider','gemini');monkeypatch.setattr(settings,'ai_concurrency',3)
    active=0;maximum=0;calls=[]
    async def fake(task,data,schema):
        nonlocal active,maximum
        active+=1;maximum=max(maximum,active);calls.append(data)
        await asyncio.sleep(.001);active-=1
        targets=[c for c in data['previous_comments'] if c['persona_id']!=data['persona']['id']]
        return {'stance':'conditional','comment':'Saya meminta prosedur yang bisa ditinjau.','rationale':'Kapasitas pelaksanaan perlu dipastikan.','issue':'capacity','proposal':'Uji prosedur bersama kelompok terdampak.','reply_to':targets[-1]['id'] if targets else None,'evidence_ids':[]},{}
    monkeypatch.setattr(ai,'generate',fake)
    draft=client.get('/api/drafts').json()[0]
    result=asyncio.run(social.simulate(draft,'Alternatif pelayanan yang inklusif.',{},[],lambda *args:None,{'persona_count':32,'rounds':5,'crowd_size':0}))
    assert maximum<=3 and len(calls)==320
    assert all(len(c['previous_comments'])<=96 for c in calls)
    assert result['provider']['mode']=='live'

@pytest.mark.parametrize('bad',[{'persona_count':64},{'rounds':6},{'rounds':1}])
def test_network_options_reject_unsupported_sizes(bad):
    with pytest.raises(ValueError):SocialOptions.model_validate(bad)

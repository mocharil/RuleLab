from __future__ import annotations
import base64,io,json,re,uuid
from contextlib import asynccontextmanager
from pathlib import Path
import fitz,httpx
from docx import Document
from fastapi import FastAPI,Request,HTTPException,Query
from fastapi.responses import FileResponse,Response,JSONResponse,StreamingResponse
from fastapi.staticfiles import StaticFiles
from app import store
from app.config import ROOT,settings
from app.models import DraftInput,SaveDraft,GenerateRequest,SimulationRequest,RevisionRequest,ReviewRequest,ImportRequest,CompareRequest
from app.services import analysis,simulation,retrieval,ai,legal,agents,social,readiness,drafting
from app import security
from app.models import AnalyzeRequest
from app.services.export import export_docx
from app.services import library_chat,collection,sources

@asynccontextmanager
async def lifespan(app):
 security.validate_deployment();store.initialize();agents.initialize();collection.initialize()
 try:yield
 finally:await agents.shutdown()
app=FastAPI(title='RuleLab',version='1.3.0',lifespan=lifespan)

@app.middleware('http')
async def protections(request:Request,call_next):
 request_id=uuid.uuid4().hex
 request.state.principal={'id':'public','role':'viewer'}
 async def reject(message,status):
  agents.audit(request_id,request.state.principal['id'],request.state.principal['role'],request.method+' '+request.url.path,status)
  return JSONResponse({'detail':message,'request_id':request_id},status,headers={'X-Request-ID':request_id})
 if request.url.path.startswith('/api/') and request.url.path not in {'/api/health','/api/public/collection'}:
  principal=security.authenticate(request.headers.get('authorization',''))
  if not principal:return await reject('An access token is required.',401)
  request.state.principal=principal
  if not security.permitted(principal,request.method,request.url.path):return await reject('Your role is not allowed to perform this action.',403)
 if request.method in ('POST','PUT','PATCH','DELETE'):
  origin=request.headers.get('origin')
  allowed=[x.strip().rstrip('/') for x in settings.allowed_origins.split(',') if x.strip()]
  if origin and origin.rstrip('/') not in (allowed or [str(request.base_url).rstrip('/')]):
   return await reject('Origin not allowed.',403)
  try:length=int(request.headers.get('content-length','0') or 0)
  except ValueError:return await reject('Invalid Content-Length.',400)
  if length>15000000:return await reject('Request too large.',413)
  chunks=[];total=0
  async for chunk in request.stream():
   total+=len(chunk)
   if total>15000000:return await reject('Request too large.',413)
   chunks.append(chunk)
  # Cache the bounded body for Starlette's downstream request parser.
  request._body=b''.join(chunks)

 response=await call_next(request)
 response.headers['X-Request-ID']=request_id
 if request.url.path.startswith('/api/'):
  response.headers['Cache-Control']='no-store'
  if request.method in ('POST','PUT','PATCH','DELETE') or response.status_code>=400:agents.audit(request_id,request.state.principal['id'],request.state.principal['role'],request.method+' '+request.url.path,response.status_code)
 elif request.url.path.startswith('/static/') or request.url.path.startswith('/workspace'):
  response.headers['Cache-Control']='no-store, no-cache, must-revalidate'
 response.headers['X-Content-Type-Options']='nosniff';response.headers['Referrer-Policy']='same-origin'
 response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; object-src 'self'; frame-src 'self'; frame-ancestors 'self'"
 return response

@app.exception_handler(KeyError)
async def missing(request,exc):return JSONResponse({'detail':str(exc.args[0])},404)
@app.exception_handler(store.RevisionConflict)
async def conflict(request,exc):return JSONResponse({'detail':str(exc)},409)
@app.exception_handler(ai.AIError)
async def provider_error(request,exc):return JSONResponse({'detail':str(exc),'code':exc.code,'retryable':exc.retryable},502)
@app.exception_handler(httpx.HTTPError)
async def search_error(request,exc):return JSONResponse({'detail':'Search service unavailable. Check Elasticsearch, the index and the connection.'},503)
@app.exception_handler(ValueError)
async def invalid(request,exc):return JSONResponse({'detail':str(exc)},400)

@app.get('/api/health')
def health():return {'status':'ok','version':'1.3.0'}
@app.get('/api/legal/guide')
def legal_guide():return legal.rulebook()
@app.get('/api/legal/changes')
def legal_changes():return {'current_law_verified':False,'candidates':legal.change_candidates()}
@app.get('/api/status')
def status(request:Request):
 regs=store.regulations()
 draft_items=store.list_drafts()
 active_draft=store.get_draft(draft_items[0]['id']) if draft_items else None
 active_personas=social.get_personas_for_draft(active_draft)
 return {'principal':request.state.principal,'agent_runtime':{'enabled':True,'mode':'bounded_specialists','max_concurrency':settings.max_agent_jobs},'deployment_mode':settings.deployment_mode,'provider':settings.provider,'model':settings.model if settings.provider!='demo' else None,'search_backend':settings.search_backend,
  'regulations':len(regs),'articles':sum(len(d['articles']) for d in regs),'social_defaults':{'persona_count':24,'rounds':4},'social_personas':active_personas,'domain':social.detect_domain(active_draft),
  'credential_configured':bool(settings.api_key) if settings.provider=='gemini' else bool(settings.credentials_path or settings.project) if settings.provider=='vertex' else False,
  'corpus_notice':'Original snapshot. Not consolidated and its effective status is not verified.'}

@app.get('/api/public/collection')
def public_collection(kind:str=Query(default='',max_length=30),start:int|None=Query(default=None,ge=1945,le=2100),end:int|None=Query(default=None,ge=1945,le=2100),mode:str=Query(default='issuer',pattern='^(issuer|scope|mentions)$')):
 if start is not None and end is not None and start>end:raise ValueError('The start year cannot be after the end year.')
 return collection.overview(kind,start,end,mode,include_private=False)

@app.get('/api/data/overview')
def data_overview(kind:str=Query(default='',max_length=30),start:int|None=Query(default=None,ge=1945,le=2100),end:int|None=Query(default=None,ge=1945,le=2100),mode:str=Query(default='issuer',pattern='^(issuer|scope|mentions)$')):
 if start is not None and end is not None and start>end:raise ValueError('The start year cannot be after the end year.')
 return collection.overview(kind,start,end,mode)

@app.post('/api/data/intake',status_code=201)
def source_intake(body:collection.IntakeRequest):return collection.add_intake(body)

@app.get('/api/data/intake/{id}')
def source_intake_detail(id:str):return collection.get_intake(id)

@app.get('/api/data/intake/{id}/raw')
def source_intake_raw(id:str):
 item=collection.get_intake(id)
 with store.db() as c:row=c.execute('SELECT content FROM source_intake WHERE id=?',(id,)).fetchone()
 if not row['content']:raise HTTPException(404,'This link has no uploaded file yet.')
 return Response(base64.b64decode(row['content']),media_type='application/octet-stream',headers={'Content-Disposition':'attachment; filename="source'+Path(item['filename']).suffix+'"'})
@app.get('/api/readiness')
def readiness_status():return readiness.inspect()
@app.post('/api/readiness/probe')
async def readiness_probe(request:Request):
 if request.state.principal['role']!='admin':raise HTTPException(403,'The integration test is only available to workspace admins.')
 return await readiness.probe()
@app.get('/api/social/personas')
def social_personas(draft_id:str,count:int=Query(default=24)):
 return {'domain':social.detect_domain(store.get_draft(draft_id)),'personas':social.get_personas_for_draft(store.get_draft(draft_id),count)}
@app.get('/api/drafting/templates')
def drafting_templates():return drafting.catalog()
@app.post('/api/drafting/preview')
def drafting_preview(body:drafting.TemplateRequest):return drafting.build(body)
@app.post('/api/drafting/create',status_code=201)
def drafting_create(body:drafting.TemplateRequest):
 result=drafting.build(body)
 return {**result,'draft':store.create_draft(result['draft'])}
@app.get('/api/drafts')
def drafts():
 items=store.list_drafts()
 return [{**d, 'domain': social.detect_domain(d), 'personas': social.get_personas_for_draft(d)} for d in items]
@app.post('/api/drafts',status_code=201)
def create(body:DraftInput):return store.create_draft(body.model_dump())
@app.get('/api/drafts/{id}')
def get_draft(id:str):
 draft=store.get_draft(id)
 return {**draft, 'domain': social.detect_domain(draft), 'personas': social.get_personas_for_draft(draft)}
@app.put('/api/drafts/{id}')
def save(id:str,body:SaveDraft):return store.save_draft(id,body.model_dump(exclude={'expected_revision'}),body.expected_revision)
@app.get('/api/drafts/{id}/history')
def history(id:str):return store.history(id)
@app.get('/api/drafts/{id}/runs')
def runs(id:str):store.get_draft(id);return store.runs(id)
@app.post('/api/drafts/{id}/analyze')
async def analyze(id:str,body:AnalyzeRequest|None=None):
 draft=store.get_draft(id);result=await analysis.analyze(draft,as_of=body.as_of_date.isoformat() if body and body.as_of_date else None);return store.add_run(draft,'analysis',result)
@app.post('/api/drafts/{id}/simulate')
async def simulate(id:str,body:SimulationRequest):
 draft=store.get_draft(id);result=await simulation.simulate(draft,body.alternative_text,body.assumptions.model_dump());return store.add_run(draft,'simulation',result)
@app.post('/api/drafts/{id}/apply-revision')
def apply(id:str,body:RevisionRequest):
 draft=store.get_draft(id);run=store.get_run(body.run_id)
 if run['draft_id']!=id or run['kind']!='analysis':raise HTTPException(400,'This analysis does not belong to this draft.')
 if run['revision']!=draft['revision'] or body.expected_revision!=draft['revision']:raise store.RevisionConflict('The analysis comes from an older version. Run the check again.')
 findings=run['payload']['findings']
 if body.finding_index>=len(findings):raise HTTPException(400,'Finding not found.')
 finding=findings[body.finding_index]
 for a in draft['articles']:
  if a['number']==finding['article_number']:a['text']=finding['proposed_text']
 payload=DraftInput.model_validate(draft).model_dump()
 return store.save_draft(id,payload,body.expected_revision)
@app.post('/api/runs/{id}/review')
def review(id:str,body:ReviewRequest):return store.review_run(id,body.status,body.notes)
@app.post('/api/generate')
async def generate(body:GenerateRequest):
 result=await analysis.generate_draft(body.brief,body.regulation_type);draft=store.create_draft(result['draft']);return {**result,'draft':draft}
@app.get('/api/corpus')
def corpus():return [{k:v for k,v in d.items() if k!='articles'} for d in store.regulations()]
@app.post('/api/gallery/chat')
async def gallery_chat(body:library_chat.ChatRequest):return await library_chat.answer(body)
@app.get('/api/corpus/{id}')
def regulation(id:str):
 result=next((d for d in store.regulations() if d['id']==id),None)
 if not result:raise HTTPException(404,'Regulation not found.')
 return result
@app.get('/api/corpus/{id}/raw')
def raw(id:str,inline:bool=Query(default=False)):
 d=regulation(id);rel=d['provenance']['raw_file']
 if not sources.safe(rel) or not sources.exists(rel):raise HTTPException(404,'Source document not available.')
 name=Path(rel).name;pdf=name.lower().endswith('.pdf')
 media='application/pdf' if pdf else 'text/html; charset=utf-8'
 disposition=('inline' if inline else 'attachment')+f'; filename="{name}"'
 try:return StreamingResponse(sources.stream(rel),media_type=media,headers={'Content-Disposition':disposition})
 except FileNotFoundError:raise HTTPException(404,'Source document not available.')
@app.get('/api/corpus/{id}/json')
def enriched(id:str):return Response(json.dumps(regulation(id),ensure_ascii=False,indent=2),media_type='application/json',headers={'Content-Disposition':f'attachment; filename="{id}.json"'})
@app.get('/api/search')
async def search(q:str=Query(min_length=2,max_length=4000),limit:int=Query(default=12,ge=1,le=50),regulation_id:str|None=None):return {'query':q,'results':await retrieval.search(q,limit,regulation_id)}
@app.post('/api/compare')
async def compare(body:CompareRequest):
 left=regulation(body.left_id);right=regulation(body.right_id)
 la=next((a for a in left['articles'] if a['number']==body.left_article),None);ra=next((a for a in right['articles'] if a['number']==body.right_article),None)
 if not la or not ra:raise HTTPException(404,'Article not found.')
 ld=store.article(la['id']);rd=store.article(ra['id'])
 draft={'id':'comparison','revision':1,'title':left['metadata']['title'],'regulation_type':left['metadata']['type'],'articles':[{'number':la['number'],'text':la['text']}]}
 if settings.provider=='demo':
  output={'summary':'Source comparison is ready for review. Demo mode does not assess conflicts between two real regulations automatically.',
   'findings':[],'limitations':analysis.LIMITATIONS,'provider':{'mode':'demo','provider':'side_by_side'},'evidence':[ld,rd],
   'hierarchy_relation':analysis.relation(ld['regulation_type'],rd['regulation_type'])}
 else:output=await analysis.analyze(draft,[ld,rd]);output['hierarchy_relation']=analysis.relation(ld['regulation_type'],rd['regulation_type'])
 return {**output,'left':ld,'right':rd}


def parse_draft_text(text,title='Imported draft',regulation_type='PP'):
 matches=list(re.finditer(r'(?m)^[ \t]*Pasal[ \t]+(\d+[A-Z]?|I|II)[ \t]*$',text))
 if not matches:raise ValueError('The document must contain Pasal 1, Pasal 2, etc. headings on their own lines.')
 articles=[]
 for i,m in enumerate(matches):
  end=matches[i+1].start() if i+1<len(matches) else len(text)
  body=text[m.end():end].strip()
  if body:articles.append({'number':m.group(1),'text':body})
 return DraftInput(title=title,regulation_type=regulation_type,preamble=text[:matches[0].start()].strip(),articles=articles).model_dump()

@app.post('/api/import')
def import_draft(body:ImportRequest):
 try:data=base64.b64decode(body.content_base64,validate=True)
 except Exception:raise HTTPException(400,'Invalid Base64.')
 if len(data)>10*1024*1024:raise HTTPException(413,'Maximum 10MB.')
 suffix=Path(body.filename).suffix.lower()
 if suffix=='.json':
  try:payload=DraftInput.model_validate_json(data).model_dump()
  except ValueError:raise HTTPException(400,'JSON must match the DraftInput schema (title, regulation_type, articles).')
 else:
  if suffix=='.txt':
   try:text=data.decode('utf-8-sig')
   except UnicodeDecodeError:raise HTTPException(400,'TXT must be UTF-8.')
  elif suffix=='.docx':
   # Reject zip bombs before the Office parser expands contents.
   import zipfile
   try:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
     if sum(i.file_size for i in z.infolist())>30*1024*1024:raise HTTPException(413,'DOCX is too large after extraction.')
    text='\n'.join(p.text for p in Document(io.BytesIO(data)).paragraphs)
   except HTTPException:raise
   except Exception:raise HTTPException(400,'The DOCX could not be read.')
  elif suffix=='.pdf':
   try:
    with fitz.open(stream=data,filetype='pdf') as pdf:
     if len(pdf)>150:raise HTTPException(413,'Maximum 150 pages.')
     text='\n'.join(p.get_text() for p in pdf)
    if len(text.strip())<30:raise HTTPException(400,'Scanned PDFs need OCR; import the OCR result as TXT/DOCX first.')
   except HTTPException:raise
   except Exception:raise HTTPException(400,'The PDF could not be read.')
  else:raise HTTPException(400,'Use TXT, JSON, DOCX or text PDF.')
  payload=parse_draft_text(text,Path(body.filename).stem)
 return store.create_draft(payload)

@app.get('/api/drafts/{id}/export')
def export(id:str,format:str=Query(default='docx',pattern='^(docx|json|report)$')):
 draft=store.get_draft(id)
 if format=='json':data=json.dumps(draft,ensure_ascii=False,indent=2).encode();mime='application/json';name='rulelab-draft.json'
 else:data=export_docx(draft,store.runs(id),format=='report');mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document';name='rulelab-review.docx' if format=='report' else 'rulelab-draft.docx'
 return Response(data,media_type=mime,headers={'Content-Disposition':f'attachment; filename="{name}"'})
app.mount('/static',StaticFiles(directory=ROOT/'app/static'),name='static')
@app.get('/')
def index():return FileResponse(ROOT/'app/static/landing.html')
@app.get('/guide')
def user_guide():return FileResponse(ROOT/'app/static/guide.html')
@app.get('/guide/download')
def user_guide_download():return FileResponse(ROOT/'docs/USER_GUIDE.md',filename='RuleLab-User-Guide.md',media_type='text/markdown; charset=utf-8')

@app.get('/workspace')
def workspace():return FileResponse(ROOT/'app/static/index.html')


@app.post('/api/drafts/{id}/agent-jobs',status_code=202)
async def start_agent(id:str,body:agents.JobRequest,request:Request):
 job_id,created=agents.create(id,body,request.state.principal['id'],request.headers.get('idempotency-key',''))
 if created:agents.schedule(job_id)
 return agents.get(job_id)

@app.get('/api/drafts/{id}/agent-jobs')
def draft_agent_jobs(id:str):return agents.list_jobs(id)

@app.get('/api/agent-jobs/{id}')
def agent_job(id:str):return agents.get(id)

@app.post('/api/agent-jobs/{id}/cancel')
async def cancel_agent(id:str,request:Request):
 try:return agents.cancel(id,request.state.principal)
 except PermissionError as exc:raise HTTPException(403,str(exc))

@app.get('/api/audit')
def audit_log(request:Request,limit:int=Query(default=100,ge=1,le=500)):
 if request.state.principal['role'] not in {'admin','reviewer'}:raise HTTPException(403,'Audit requires the admin or reviewer role.')
 return {'events':agents.audit_rows(limit)}

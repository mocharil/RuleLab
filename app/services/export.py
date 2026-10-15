from io import BytesIO
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def export_docx(draft,runs,report=False):
 document=Document();s=document.sections[0];s.top_margin=Cm(2.5);s.bottom_margin=Cm(2.5);s.left_margin=Cm(3);s.right_margin=Cm(2.5)
 style=document.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(12);style.paragraph_format.space_after=Pt(7)
 header=s.header.paragraphs[0];header.text='RULELAB  /  DRAFT FOR REVIEW';header.style='Caption'
 heading=draft.get('document_heading') or ('RANCANGAN '+draft['regulation_type']+'\nTENTANG\n'+draft['title'].upper())
 for line in heading.splitlines():
  if not line.strip():continue
  p=document.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].bold=True
 cap=document.add_paragraph(f"Draft version {draft['revision']} • Not enacted • Needs legal expert review",style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
 if draft.get('preamble'):
  for raw_line in draft['preamble'].splitlines():
   line=raw_line.strip()
   if not line:
    continue
   p=document.add_paragraph(line)
   if any(k in line.upper() for k in ['DENGAN RAHMAT TUHAN', 'PRESIDEN REPUBLIK', 'GUBERNUR ', 'BUPATI ', 'WALIKOTA ', 'MEMUTUSKAN:', 'DENGAN PERSETUJUAN BERSAMA', 'DEWAN PERWAKILAN RAKYAT']):
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if p.runs: p.runs[0].bold=True
   elif line.startswith(('Menimbang:', 'Mengingat:', 'Menetapkan:')):
    if p.runs: p.runs[0].bold=True
 for a in draft['articles']:
  if a.get('chapter'):
   for line in a['chapter'].splitlines():
    if not line.strip():continue
    p=document.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].bold=True
  h=document.add_paragraph('Pasal '+a['number']);h.alignment=WD_ALIGN_PARAGRAPH.CENTER;h.runs[0].bold=True
  for line in a['text'].splitlines():
   t=line.strip()
   if t:
    p=document.add_paragraph(t)
    if t.startswith(('(', 'a.', 'b.', 'c.', 'd.', 'e.', 'f.', '1.', '2.', '3.')):
     p.paragraph_format.left_indent=Cm(0.7)
 for line in draft.get('closing','').splitlines():
  p=document.add_paragraph(line)
 for key,title in [('explanation','PENJELASAN'),('attachments','LAMPIRAN')]:
  if draft.get(key):
   document.add_page_break();document.add_heading(title,1)
   for line in draft[key].splitlines():document.add_paragraph(line)
 if report:
  document.add_page_break();document.add_heading('RuleLab review annex',0)
  document.add_paragraph('AI report to support review, not a determination of legal validity. Only analyses for this draft version are attached.')
  for run in runs:
   if run['revision']!=draft['revision']:continue
   p=run['payload'];document.add_heading('Conflict analysis' if run['kind']=='analysis' else 'Synthetic public forum' if run['kind']=='social' else 'Impact simulation',1)
   document.add_paragraph(f"Provider: {p['provider']['provider']} | Mode: {p['provider']['mode']} | Review: {run['review_status']}")
   if run.get('review_notes'):document.add_paragraph('Reviewer notes: '+run['review_notes'])
   if run['kind']=='analysis':
    document.add_paragraph(p['summary'])
    if p.get('legal_context'):
     ctx=p['legal_context'];document.add_paragraph('Guidelines: '+ctx['rulebook_version']+' | Analysis date: '+ctx['as_of_date'])
     document.add_paragraph(ctx['notice'])
    for f in p['findings']:
     document.add_heading('Pasal '+f['article_number']+': '+f['title'],2);document.add_paragraph(f['reasoning'])
     document.add_paragraph('Alternative interpretation: '+f['counterargument'])
     document.add_paragraph('Guideline rules: '+', '.join(f.get('legal_rule_ids',[])))
     for warning in f.get('validation',{}).get('warnings',[]):document.add_paragraph('Validation: '+warning)
     for e in f['evidence']:
      document.add_paragraph(e['regulation_label']+' Pasal '+e['number']+' | '+e['source']['url']+' | Page '+', '.join(map(str,e['pages'])))
     document.add_paragraph('Proposal: '+f['proposed_text'])
   elif run['kind']=='social':
    document.add_paragraph('Assumptions: '+str(p['assumptions']))
    for scenario in ['A','B']:
     document.add_heading('Scenario '+scenario,2)
     for comment in p['scenarios'][scenario]['comments']:
      document.add_paragraph(comment['name']+' | Round '+str(comment['round'])+' | '+comment['stance'])
      document.add_paragraph(comment['comment'])
      document.add_paragraph('Rationale: '+comment['rationale'])
      document.add_paragraph('Reply to: '+str(comment['reply_to']))
      document.add_paragraph('Proposal: '+comment['proposal'])
   else:
    document.add_paragraph('Assumptions: '+str(p['assumptions']));document.add_paragraph('Result A: '+str(p['A']));document.add_paragraph('Result B: '+str(p['B']))
    for person in p['stakeholders']:
     document.add_heading(person['stakeholder'],2)
     for risk in person['risks']:document.add_paragraph(risk,style='List Bullet')
   for limitation in p['limitations']:document.add_paragraph(limitation,style='List Bullet')
 footer=s.footer.paragraphs[0];footer.text='Draft for deliberation. It has no status as an enacted regulation.';footer.style='Caption'
 stream=BytesIO();document.save(stream);return stream.getvalue()
